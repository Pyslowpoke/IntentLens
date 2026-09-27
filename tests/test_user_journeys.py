import io
import json
import pandas as pd
import pytest
from openpyxl import Workbook
from fastapi.testclient import TestClient
from services.api.main import app
from services.api import store
from services.api.providers import ProviderAdapter, Proposals
from services.api.render import render
from services.api.contracts import ChartSpec


@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.setenv('WORKBENCH_DATA',str(tmp_path/'data'))
    return TestClient(app)


def workbook():
    wb=Workbook();wb.active.title='空白说明'
    ws=wb.create_sheet('真实业务');ws.append(['报告标题']);ws.append(['地区','金额']);ws.append(['East',10]);ws.append([None,20])
    wb.create_sheet('隐藏数据').sheet_state='hidden'
    out=io.BytesIO();wb.save(out);return out.getvalue()


def test_sheet_discovery_before_parsing_and_exact_import(client):
    raw=workbook()
    catalog=client.post('/imports/sheets',files={'file':('multi.xlsx',raw)}).json()['sheets']
    assert [s['name'] for s in catalog]==['空白说明','真实业务','隐藏数据']
    assert catalog[2]['hidden']
    assert client.post('/imports/preview',files={'file':('multi.xlsx',raw)}).status_code==422
    preview=client.post('/imports/preview',files={'file':('multi.xlsx',raw)},data={'sheet':'真实业务','header':2}).json()
    assert preview['names']==['地区','金额'] and preview['rows']==2
    project=client.post('/projects',json={'name':'sheet-test'}).json()
    ds=client.post(f'/projects/{project["id"]}/import',json={'upload_id':preview['id']}).json()
    assert ds['transforms'][0]['options']['sheet']=='真实业务'
    df=pd.read_parquet(store.root()/ds['snapshot'])
    assert df.f2.sum()==30


def test_model_discussion_restores_history_and_resets_on_dataset_change(client,monkeypatch):
    seen=[]
    def request(self,prompt,contract=Proposals):
        seen.append(json.loads(prompt))
        return Proposals(analysis='Which metric should be compared?',questions=['Revenue or count?'])
    monkeypatch.setattr(ProviderAdapter,'request',request)
    p=client.post('/projects',json={'name':'conversation-test'}).json();url='/projects/'+p['id']
    client.post(url+'/sample/sales',json={})
    a=client.post(url+'/recommend',json={'goal':'帮我看看','mode':'model'}).json()
    assert not a['proposals'] and len(a['messages'])==2
    client.post(url+'/recommend',json={'goal':'比较销售额','mode':'model'})
    assert seen[-1]['conversation'][0]['content']=='帮我看看'
    assert len(client.get(url+'/context').json()['messages'])==4
    client.post(url+'/sample/product',json={})
    assert client.get(url+'/context').json()['messages']==[]
    client.post(url+'/recommend',json={'goal':'看转化','mode':'model'})
    assert seen[-1]['conversation']==[]


def test_matplotlib_nullable_categories_renders_real_nonblank_png(tmp_path):
    result={'records':[{'region':'A','value':10},{'region':None,'value':20}], 'columns':['region','value'],'labels':{}}
    artifacts=render(result,ChartSpec(engine='matplotlib',kind='bar'),tmp_path)
    from PIL import Image,ImageStat
    with Image.open(tmp_path/artifacts['png']) as image:
        assert image.size==(1000,560)
        assert sum(ImageStat.Stat(image.convert('RGB')).var)>100


def test_model_field_names_resolve_only_when_unambiguous(client,monkeypatch):
    from services.api.providers import recommend
    p=client.post('/projects',json={'name':'field-test'}).json()
    dataset=client.post(f'/projects/{p["id"]}/paste',json={'text':'地区\t金额\nEast\t10'}).json()
    monkeypatch.setattr(ProviderAdapter,'request',lambda *args:Proposals.model_validate({'proposals':[{'title':'Compare','explanation':'Test','plan':{'group_by':['地区'],'metric':'金额'},'spec':{'mapping':{'x':'地区','y':'金额'}}}]}))
    result=recommend(dataset,'compare','model')
    assert result['proposals'][0]['plan']['group_by']==['f1']
    assert result['proposals'][0]['spec']['mapping']=={'x':'f1','y':'value'}


def test_artifact_download_is_attachment_but_preview_stays_inline(client):
    folder=store.root()/'artifacts'/'download-test';folder.mkdir(parents=True)
    (folder/'chart.png').write_bytes(b'\x89PNG\r\n\x1a\n')
    store.put('revision',{'id':'download-test','artifacts':{'png':'chart.png'}})
    preview=client.get('/revisions/download-test/artifact/png')
    assert 'content-disposition' not in preview.headers
    download=client.get('/revisions/download-test/artifact/png?download=true')
    assert download.status_code==200
    assert download.headers['content-disposition'].startswith('attachment;')
    assert download.content==preview.content


@pytest.mark.parametrize('percent',[False,True])
def test_horizontal_altair_keeps_category_labels(tmp_path,percent):
    result={'records':[{'person':'示例甲','value':.75},{'person':'示例乙','value':.43}], 'columns':['person','value'],'labels':{'person':'人员','value':'占比' if percent else '金额'},'value_format':'percent' if percent else 'number'}
    render(result,ChartSpec(engine='altair',kind='bar',mapping={'x':'value','y':'person'}),tmp_path)
    spec=json.loads((tmp_path/'vega.json').read_text(encoding='utf-8'))
    assert spec['encoding']['y']['type']=='nominal'
    assert 'format' not in spec['encoding']['y']['axis']
    assert spec['encoding']['x']['type']=='quantitative'
    if percent: assert spec['encoding']['x']['axis']['format']=='.1%'
    svg=(tmp_path/'chart.svg').read_text(encoding='utf-8')
    assert '示例甲' in svg and '示例乙' in svg
    assert 'NaNundefined' not in svg
