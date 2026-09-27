import csv
import hashlib
import io
import json
import os
from pathlib import Path
import pandas as pd
from .contracts import DatasetSnapshot, DatasetField
from .store import root, uid, put


def workbook_sheets(raw):
    import zipfile
    import openpyxl
    if len(raw)>int(os.getenv("MAX_UPLOAD_MB","50"))*1024*1024: raise ValueError("File exceeds upload limit")
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        if sum(i.file_size for i in archive.infolist())>int(os.getenv("MAX_UPLOAD_MB","50"))*1024*1024*20: raise ValueError("Expanded workbook too large")
    wb=openpyxl.load_workbook(io.BytesIO(raw),data_only=True,read_only=True)
    try:
        return [{"name":ws.title,"rows":ws.max_row or 0,"columns":ws.max_column or 0,"hidden":ws.sheet_state!="visible","preview":[[str(v)[:120] if v is not None else "" for v in row] for row in ws.iter_rows(max_row=5,max_col=min(ws.max_column or 1,8),values_only=True)]} for ws in wb.worksheets]
    finally: wb.close()


def parse(raw: bytes, name: str, sheet=None, header=1, delimiter=None, encoding="utf-8-sig"):
    if len(raw) > int(os.getenv("MAX_UPLOAD_MB", "50")) * 1024 * 1024:
        raise ValueError("File exceeds configured upload limit / 文件过大")
    if header < 1:
        raise ValueError("Header row starts at 1")
    warnings = []
    sheets = []
    if name.lower().endswith(".xlsx"):
        import zipfile
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            expanded=sum(info.file_size for info in archive.infolist())
            if expanded > int(os.getenv("MAX_UPLOAD_MB","50"))*1024*1024*20:
                raise ValueError("Expanded workbook exceeds configured safety limit / Excel 解压后体积过大")
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
        formulas = openpyxl.load_workbook(io.BytesIO(raw), data_only=False, read_only=True)
        sheets = wb.sheetnames
        chosen = sheet or sheets[0]
        if chosen not in sheets:
            raise ValueError("Worksheet not found")
        ws = wb[chosen]
        if ws.max_column > 1000:
            raise ValueError("Workbook exceeds 1000 columns; select a narrower source")
        if ws.max_row > int(os.getenv("MAX_ROWS", "1000000")) + header:
            raise ValueError("Worksheet exceeds configured row limit")
        rows = list(ws.iter_rows(min_row=header, values_only=True))
        for row in formulas[chosen].iter_rows(min_row=header + 1):
            for cell in row:
                if cell.data_type == "f" and rows[cell.row-header][cell.column-1] is None:
                    warnings.append(f"Formula cache missing at {cell.coordinate}; open and recalculate in Excel / 公式缓存缺失")
        wb.close()
        formulas.close()
    elif name.lower().endswith((".csv", ".tsv", ".txt")):
        try:
            text = raw.decode(encoding, errors="strict")
        except (UnicodeError, LookupError) as e:
            raise ValueError("Encoding failed; choose UTF-8 or GB18030 / 请修正编码") from e
        delim = delimiter or ("\t" if name.lower().endswith(".tsv") else ",")
        if len(delim) != 1:
            raise ValueError("Delimiter must be one character")
        try:
            rows = list(csv.reader(io.StringIO(text), delimiter=delim, strict=True))[header - 1:]
        except csv.Error as e:
            raise ValueError(f"CSV parsing failed; check quotes and delimiter / CSV 格式错误: {e}") from e
    else:
        raise ValueError("Supported files: .xlsx .csv .tsv / 不执行宏")
    if not rows or not rows[0]:
        raise ValueError("Empty header / 表头为空")
    names = [str(v) if v is not None else "Unnamed" for v in rows[0]]
    values = rows[1:]
    if len(values) > int(os.getenv("MAX_ROWS", "1000000")):
        raise ValueError("Row limit exceeded; no rows were imported")
    for i, row in enumerate(values):
        if len(row) != len(names):
            raise ValueError(f"Row {i + header + 1}: expected {len(names)} columns, got {len(row)}; fix delimiter / 列数不一致")
    if len(names) != len(set(names)):
        warnings.append("Duplicate headers retained with stable IDs / 重复表头已保留并分配稳定 ID")
    df = pd.DataFrame(values, columns=[f"f{i+1}" for i in range(len(names))]).replace("", None)
    return df, names, warnings, sheets


def snapshot(df, names, source, raw=b"", warnings=None, overrides=None, transforms=None):
    df = df.copy()
    warnings = list(warnings or [])
    fields, profile = [], {}
    overrides = overrides or {}
    for id, name in zip(df.columns, names):
        s = df[id]
        nonnull = s.dropna()
        numeric = pd.to_numeric(s, errors="coerce")
        dtype = "string"
        isdate = pd.api.types.is_datetime64_any_dtype(s)
        if len(nonnull) and numeric.notna().sum() == len(nonnull) and not isdate:
            df[id], dtype = numeric, "number"
        elif len(nonnull) and (pd.api.types.is_datetime64_any_dtype(s) or all(isinstance(v, (pd.Timestamp, __import__('datetime').datetime)) or __import__('re').match(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}", str(v)) for v in nonnull)):
            dates = pd.to_datetime(s, errors="coerce")
            if dates.notna().sum() == len(nonnull):
                df[id], dtype = dates, "date"
        elif 0 < numeric.notna().sum() < len(nonnull):
            warnings.append(f"{name}: mixed numeric/text values; kept as text / 混合数字文本，未丢弃非法值")
        requested = overrides.get(id, {})
        target = requested.get("dtype", dtype)
        if target == "number":
            converted = pd.to_numeric(s, errors="coerce")
            if (s.notna() & converted.isna()).any():
                raise ValueError(f"{name}: invalid numeric values / 非法数字: {s[s.notna() & converted.isna()].head(3).tolist()}")
            import numpy as np
            if not np.isfinite(converted.dropna()).all():
                raise ValueError(f"{name}: infinite numeric values are not supported / 不支持无限值")
            df[id] = converted
        elif target == "date":
            converted = pd.to_datetime(s, errors="coerce")
            if (s.notna() & converted.isna()).any():
                raise ValueError(f"{name}: invalid date values")
            df[id] = converted
        elif target == "string":
            df[id] = s.astype("string")
        elif target == "boolean":
            lookup = {"true": True, "false": False, "1": True, "0": False}
            converted = s.map(lambda x: lookup.get(str(x).lower()) if pd.notna(x) else None)
            if (s.notna() & converted.isna()).any():
                raise ValueError(f"{name}: boolean must be true/false/1/0")
            df[id] = converted.astype("boolean")
        fields.append(DatasetField(id=id, name=name, dtype=target, meaning=requested.get("meaning", ""), nulls=int(df[id].isna().sum()), unique=int(df[id].nunique())))
        profile[id] = {"nulls": int(df[id].isna().sum()), "unique": int(df[id].nunique())}
        if target in ("number", "date") and df[id].notna().any():
            profile[id].update(min=str(df[id].min()), max=str(df[id].max()))
    profile["duplicate_rows"] = int(df.duplicated().sum())
    id = uid()
    path = root() / "snapshots" / f"{id}.parquet"
    df.to_parquet(path, index=False)
    rawpath = root() / "raw" / id
    rawpath.write_bytes(raw)
    obj = DatasetSnapshot(id=id, source=source, fields=fields, rows=len(df), checksum=hashlib.sha256(path.read_bytes()).hexdigest(), snapshot=str(path.relative_to(root())), raw=str(rawpath.relative_to(root())), warnings=warnings, transforms=transforms or [], profile=profile).model_dump()
    put("dataset", obj)
    return obj


def records(df):
    return json.loads(df.to_json(orient="records", date_format="iso", force_ascii=False))
